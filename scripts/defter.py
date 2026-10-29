#!/usr/bin/env python3
"""Portföy Defteri: 4 hesabın (Midas-TR, Info-TR ₺ · Midas-USD, IBKR $) günlük kaydı.

Veri: docs/defter/data/defter.json (PORTFOLIO_KEY ile şifreli; sayfa: docs/defter/index.html).
Kayıt alanları (gün → hesap):
  v  performans tabanı. TL hesap: toplam varlık − yedek akçe ("portföye ayırdığım değer").
     USD hesap: hesap değeri (hisse + nakit).
  h  hisse (pozisyon) piyasa değeri (varsa)
  c  günlük performans değişimi (para giriş/çıkışı ve yedek akçe değişimi ayıklanmış)
  f  o günkü para akışı (yatırma +, çekme −)       r  o günkü yedek akçe (TL)
  s  kaynak: x=Excel, a=otomatik (repo), m=elle    n  not
Gün düzeyi: xu (XU100 kapanış), sp (S&P 500 kapanış), n (not).

Komutlar:
  kutu-anahtar  Claude'un repoya şifreli veri bırakabilmesi için RSA anahtar çifti üretir
                (özel anahtar PORTFOLIO_KEY ile şifreli saklanır; açık anahtar defter/box_pub.pem).
  ice-aktar     defter/inbox/*.box.json kutularını açar ve deftere işler (sonra kutular silinir).
  ibkr          IBKR Flex raporundan (salt okunur token) günlük net değer + para hareketleri.
  oto           Belirtilen gün (vars. bugün, UTC) için Midas-TR / Midas-USD değerini repodan, endeksleri piyasadan yazar.
  gir           Elle kayıt (iş akışı formundan): IN_* ortam değişkenleri.
  excel         Excel dosyası üretip Telegram'dan gönderir (+ haftalık özet).
"""
import argparse
import base64
import json
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common  # noqa: E402

F_DEFTER = "docs/defter/data/defter.json"
F_BOX_PRIV = "defter/box_key.json"
F_BOX_PUB = "defter/box_pub.pem"
INBOX = "defter/inbox"
F_INFO = "defter/info_portfolio.json"
common.SECRET_FILES.update({F_DEFTER, F_BOX_PRIV, F_INFO})

ACCOUNTS = {
    "midas_tr": {"name": "Midas-TR", "cur": "TRY", "bench": "xu", "reserve": 90000, "src": "bist/history.json"},
    "info_tr": {"name": "Info-TR", "cur": "TRY", "bench": "xu", "reserve": 25000, "src": "info"},
    "midas_usd": {"name": "Midas-USD", "cur": "USD", "bench": "sp", "reserve": 0, "src": "history.json",
                  "start": "2025-10-14"},
    "ibkr": {"name": "IBKR", "cur": "USD", "bench": "sp", "reserve": 0, "src": None, "start": "2026-04-17"},
}
ALIASES = {"midas_tr": "midas_tr", "midastr": "midas_tr", "midas-tr": "midas_tr", "info": "info_tr",
           "info_tr": "info_tr", "info-tr": "info_tr", "midas_usd": "midas_usd", "midasusd": "midas_usd",
           "midas-usd": "midas_usd", "midas": "midas_usd", "ibkr": "ibkr"}
DIAG = []   # kaynak hataları (tutar içermez; status.json'a yazılır)
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/124 Safari/537.36"}
GUN = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]


# ------------------------------------------------------------------ veri
def load():
    d = common.load_json(F_DEFTER) or {}
    d.setdefault("v", 1)
    acc = d.setdefault("accounts", {})
    for k, a in ACCOUNTS.items():
        cur = acc.setdefault(k, {})
        for kk in ("name", "cur", "bench", "start"):
            if kk in a:
                cur[kk] = a[kk]
        cur.setdefault("reserve", a["reserve"])
    d.setdefault("days", [])
    return d


def save(d):
    d["days"].sort(key=lambda x: x["d"])
    d["updated"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    if not common.vault_key():
        raise common.VaultError("PORTFOLIO_KEY yok: defter düz yazılmaz")
    try:   # mini uygulama "İşlemler" sekmesi (defter.json içinde, şifreli)
        d["islem"] = trade_snapshot(d)
    except Exception as ex:
        DIAG.append(f"islem: {type(ex).__name__}")
    common.save_json(F_DEFTER, d)
    # Tutar içermeyen durum dosyası (izleme / hata ayıklama için)
    st = {"updated": d["updated"], "days": len(d["days"]),
          "first": d["days"][0]["d"] if d["days"] else None, "last": d["days"][-1]["d"] if d["days"] else None,
          "with_xu": sum(1 for x in d["days"] if "xu" in x), "with_sp": sum(1 for x in d["days"] if "sp" in x),
          "last_by_account": {a: next((x["d"] for x in reversed(d["days"]) if a in x.get("a", {})), None)
                              for a in ACCOUNTS}, "diag": DIAG[-10:]}
    common.save_json("docs/defter/data/status.json", st)


def day(d, ds, create=True):
    for x in d["days"]:
        if x["d"] == ds:
            return x
    if not create:
        return None
    x = {"d": ds, "a": {}}
    d["days"].append(x)
    d["days"].sort(key=lambda y: y["d"])
    return x


def series(d, acc):
    return [(x["d"], x["a"][acc]) for x in d["days"] if acc in x.get("a", {})]


def reserve_of(d, acc, e):
    return e.get("r", d["accounts"][acc].get("reserve", 0)) if d["accounts"][acc]["cur"] == "TRY" else 0


def recompute(d, acc):
    """Otomatik/elle kayıtların değişimini (c) bir önceki kayda göre yeniden hesaplar. Excel kayıtlarına dokunmaz."""
    prev = None
    for _, e in series(d, acc):
        if e.get("s") in ("a", "m") and prev is not None:
            tot = e["v"] + reserve_of(d, acc, e)
            ptot = prev["v"] + reserve_of(d, acc, prev)
            if prev.get("s") == "x" and "pt" in e:   # Excel→repo geçişi: tabanı repodaki önceki kayıttan al
                ptot = e["pt"]
            e["c"] = round(tot - ptot - e.get("f", 0), 2)
        prev = e


# ------------------------------------------------------------------ piyasa verisi
def yahoo_closes(symbol, rng="10d"):
    import time
    params = ({"period1": 1577836800, "period2": int(datetime.now(timezone.utc).timestamp()), "interval": "1d"}
              if rng == "max" else {"range": rng, "interval": "1d"})
    last = None
    for i, host in enumerate(("query1", "query2")):
        try:
            r = requests.get(f"https://{host}.finance.yahoo.com/v8/finance/chart/{symbol}",
                             params=params, headers=UA, timeout=30)
            r.raise_for_status()
            res = r.json()["chart"]["result"][0]
            off = res["meta"].get("gmtoffset", 0)
            out = {}
            for ts, c in zip(res["timestamp"], res["indicators"]["quote"][0]["close"]):
                if c:
                    out[datetime.fromtimestamp(ts + off, timezone.utc).date().isoformat()] = round(float(c), 2)
            return out
        except Exception as ex:
            last = ex
            time.sleep(4 * (i + 1))
    print(f"[WARN] yahoo {symbol}: {type(last).__name__}")
    DIAG.append(f"yahoo {symbol}: {type(last).__name__} {getattr(getattr(last, 'response', None), 'status_code', '')}")
    return {}


def twelvedata_closes(symbol, n=5000):
    key = os.environ.get("TWELVEDATA_API_KEY", "").strip()
    if not key:
        return {}
    try:
        r = requests.get("https://api.twelvedata.com/time_series",
                         params={"symbol": symbol, "interval": "1day", "outputsize": n, "apikey": key}, timeout=30)
        j = r.json()
        if j.get("status") != "ok":
            DIAG.append(f"twelvedata {symbol}: {str(j.get('code', ''))} {str(j.get('message', ''))[:60].replace(key, '***')}")
            return {}
        return {v["datetime"][:10]: round(float(v["close"]), 2) for v in j.get("values", [])}
    except Exception as ex:
        DIAG.append(f"twelvedata {symbol}: {type(ex).__name__}")
        return {}


def stooq_closes(symbol):
    """Yedek kaynak: Stooq günlük CSV (anahtarsız). ^spx, ^xu100"""
    try:
        r = requests.get("https://stooq.com/q/d/l/", params={"s": symbol, "i": "d"}, headers=UA, timeout=30)
        r.raise_for_status()
        out = {}
        for line in r.text.splitlines()[1:]:
            p = line.split(",")
            if len(p) >= 5 and re.fullmatch(r"\d{4}-\d{2}-\d{2}", p[0]):
                out[p[0]] = round(float(p[4]), 2)
        return out
    except Exception as ex:
        print(f"[WARN] stooq {symbol}: {type(ex).__name__}")
        DIAG.append(f"stooq {symbol}: {type(ex).__name__} {getattr(getattr(ex, 'response', None), 'status_code', '')}")
        return {}


def fred_closes(series="SP500"):
    """Yedek kaynak: FRED (St. Louis Fed) günlük S&P 500 kapanışı, ~10 yıl, anahtarsız."""
    try:
        r = requests.get("https://fred.stlouisfed.org/graph/fredgraph.csv", params={"id": series},
                         headers=UA, timeout=30)
        r.raise_for_status()
        out = {}
        for line in r.text.splitlines()[1:]:
            p = line.split(",")
            if len(p) == 2 and re.fullmatch(r"\d{4}-\d{2}-\d{2}", p[0]) and p[1] not in (".", ""):
                out[p[0]] = round(float(p[1]), 2)
        return out
    except Exception as ex:
        print(f"[WARN] fred {series}: {type(ex).__name__}")
        DIAG.append(f"fred {series}: {type(ex).__name__} {getattr(getattr(ex, 'response', None), 'status_code', '')}")
        return {}


def tv_close(symbol):
    try:
        r = requests.post("https://scanner.tradingview.com/global/scan",
                          json={"symbols": {"tickers": [symbol]}, "columns": ["close"]}, timeout=15)
        r.raise_for_status()
        rows = r.json().get("data", [])
        return round(float(rows[0]["d"][0]), 2) if rows and rows[0]["d"][0] else None
    except Exception as ex:
        print(f"[WARN] tradingview {symbol}: {type(ex).__name__}")
        DIAG.append(f"tradingview {symbol}: {type(ex).__name__} {getattr(getattr(ex, 'response', None), 'status_code', '')}")
        return None


def fill_indices(d, rng="10d", today=None):
    # S&P 500 karşılaştırması SPY fiyatıyla yapılır (endeks verisi ücretsiz kaynaklarda kısıtlı; tek kaynak = tutarlı seri)
    if d.get("sp_src") != "SPY":
        for x in d["days"]:
            x.pop("sp", None)
        d["sp_src"] = "SPY"
        rng = "max"
    xu, sp = yahoo_closes("XU100.IS", rng), yahoo_closes("SPY", rng)
    if len(sp) < 3:
        sp = {**twelvedata_closes("SPY", 5000 if rng == "max" else 15), **sp}
    if len(xu) < 3:
        xu = {**stooq_closes("^xu100"), **xu}
    if today and today not in xu:
        v = tv_close("BIST:XU100")
        if v:
            xu[today] = v
    if today and today not in sp:
        v = tv_close("AMEX:SPY")
        if v:
            sp[today] = v
    n = 0
    for x in d["days"]:
        if "xu" not in x and x["d"] in xu:
            x["xu"] = xu[x["d"]]; n += 1
        if "sp" not in x and x["d"] in sp:
            x["sp"] = sp[x["d"]]; n += 1
    return xu, sp, n


# ------------------------------------------------------------------ şifreli kutu (Claude → repo)
def cmd_box_keygen(_):
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    if not common.vault_key():
        raise common.VaultError("PORTFOLIO_KEY yok: özel anahtar düz yazılmaz")
    if common.path(F_BOX_PUB).exists():
        print("Anahtar zaten var; yeniden üretilmedi.")
        return 0
    key = rsa.generate_private_key(public_exponent=65537, key_size=3072)
    pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                            serialization.NoEncryption()).decode()
    common.save_json(F_BOX_PRIV, {"pem": pem})
    pub = key.public_key().public_bytes(serialization.Encoding.PEM,
                                        serialization.PublicFormat.SubjectPublicKeyInfo)
    common.path(F_BOX_PUB).write_bytes(pub)
    print("Kutu anahtarı üretildi.")
    return 0


def box_seal(obj, pub_pem):
    """Yerelde (Claude tarafı) kullanılır: RSA-OAEP(SHA-256) + AES-256-GCM."""
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    pub = serialization.load_pem_public_key(pub_pem)
    k, iv = os.urandom(32), os.urandom(12)
    ct = AESGCM(k).encrypt(iv, json.dumps(obj, ensure_ascii=False).encode(), None)
    ek = pub.encrypt(k, padding.OAEP(mgf=padding.MGF1(hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    b = lambda x: base64.b64encode(x).decode()  # noqa: E731
    return {"box": 1, "ek": b(ek), "iv": b(iv), "data": b(ct)}


def box_open(env):
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    priv = serialization.load_pem_private_key(common.load_json(F_BOX_PRIV)["pem"].encode(), None)
    k = priv.decrypt(base64.b64decode(env["ek"]),
                     padding.OAEP(mgf=padding.MGF1(hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    return json.loads(AESGCM(k).decrypt(base64.b64decode(env["iv"]), base64.b64decode(env["data"]), None))


def cmd_import(_):
    d = load()
    files = sorted(common.path(INBOX).glob("*.box.json")) if common.path(INBOX).exists() else []
    if not files:
        print("Gelen kutusu boş.")
        return 0
    msgs = []
    for p in files:
        item = box_open(json.loads(p.read_text()))
        kind = item.get("type")
        if kind == "days":
            n = 0
            for row in item["days"]:
                x = day(d, row["d"])
                for acc, e in row.get("a", {}).items():
                    x.setdefault("a", {})[acc] = e
                for kk in ("xu", "sp", "n"):
                    if kk in row:
                        x[kk] = row[kk]
                n += 1
            for acc, val in (item.get("reserves") or {}).items():
                d["accounts"][acc]["reserve"] = val
            msgs.append(f"📥 {n} gün işlendi ({item['days'][0]['d']} → {item['days'][-1]['d']}).")
        elif kind == "info_portfolio":
            common.save_json(F_INFO, {"cash": round(float(item["cash"]), 2), "positions": item["positions"],
                                      "transactions": [], "as_of": item.get("as_of")})
            msgs.append(f"🏦 Info-TR pozisyonları kaydedildi ({len(item['positions'])} hisse).")
        elif kind == "settings":
            for acc, kv in item.get("accounts", {}).items():
                d["accounts"].setdefault(acc, {}).update(kv)
            msgs.append("⚙️ Hesap ayarları güncellendi.")
        else:
            msgs.append(f"⚠️ Bilinmeyen kutu türü: {kind}")
        p.unlink()
    _, _, n = fill_indices(d, rng="max")
    for acc in ACCOUNTS:
        recompute(d, acc)
    msgs.append(f"📈 Endeks kapanışı eklendi: {n} hücre.")
    msgs.append(reconcile(d))
    save(d)
    common.send_telegram("📒 <b>Portföy Defteri</b>\n" + "\n".join(m for m in msgs if m))
    return 0


def reconcile(d):
    """Son Excel günü ile repodaki snapshot'ları karşılaştırır (yalnızca Telegram'a, log'a değil)."""
    lines = []
    for acc in ("midas_tr", "midas_usd"):
        a = ACCOUNTS[acc]
        ser = [(ds, e) for ds, e in series(d, acc) if e.get("s") == "x"]
        if not ser:
            continue
        ds, e = ser[-1]
        hist = common.load_json(a["src"], {}) or {}
        snap = next((s for s in hist.get("snapshots", []) if s["date"] == ds), None)
        cur = "₺" if a["cur"] == "TRY" else "$"
        if not snap:
            lines.append(f"• {a['name']} {ds}: repoda o güne ait kayıt yok.")
            continue
        res = d["accounts"][acc].get("reserve", 0)
        lines.append(f"• {a['name']} {ds}: Excel portföy {cur}{e['v']:,.0f}"
                     + (f" (hisse {cur}{e['h']:,.0f})" if e.get("h") else "")
                     + f" · Repo toplam {cur}{snap['value']:,.0f} (nakit {cur}{snap.get('cash', 0):,.0f})"
                     + (f" → toplam − yedek {cur}{snap['value'] - res:,.0f}" if res else ""))
    return ("🔎 <b>Mutabakat</b>\n" + "\n".join(lines)) if lines else ""


# ------------------------------------------------------------------ otomatik günlük kayıt
def flows_on(hist, ds):
    f = 0.0
    for t in hist.get("transactions", []):
        if t.get("date") == ds and t.get("action") in ("deposit", "withdraw"):
            f += t.get("amount", 0) * (1 if t["action"] == "deposit" else -1)
    return round(f, 2)


def cmd_auto(args):
    d = load()
    ds = args.date or datetime.now(timezone.utc).date().isoformat()
    x = day(d, ds)
    got = []
    for acc, a in ACCOUNTS.items():
        if not a["src"]:
            continue
        if a["src"] == "info":
            e0 = x["a"].get(acc)
            if e0 and e0.get("s") in ("m", "x"):
                continue
            iv = info_value()
            if not iv:
                continue
            res = d["accounts"][acc].get("reserve", 0)
            e = {"v": round(iv["total"] - res, 2), "h": round(iv["stocks"], 2), "r": res, "s": "a"}
            f = round(sum(t.get("amount", 0) * (1 if t["action"] == "deposit" else -1)
                          for t in iv["pf"].get("transactions", [])
                          if t.get("date") == ds and t.get("action") in ("deposit", "withdraw")), 2)
            if f:
                e["f"] = f
            if (e0 or {}).get("n"):
                e["n"] = e0["n"]
            if iv["missing"]:
                e["n"] = (e.get("n", "") + f" [fiyat yok: {', '.join(iv['missing'])}]").strip()
            x["a"][acc] = e
            got.append(a["name"])
            continue
        e0 = x["a"].get(acc)
        if e0 and e0.get("s") in ("m", "x"):
            continue  # elle ya da Excel'den girilmiş değere dokunma
        hist = common.load_json(a["src"], {}) or {}
        snap = next((s for s in hist.get("snapshots", []) if s["date"] == ds), None)
        if not snap:
            continue
        res = d["accounts"][acc].get("reserve", 0) if a["cur"] == "TRY" else 0
        e = {"v": round(snap["value"] - res, 2), "h": round(snap["value"] - snap.get("cash", 0), 2), "s": "a"}
        if a["cur"] == "TRY":
            e["r"] = res
        prev_snap = [s for s in hist.get("snapshots", []) if s["date"] < ds]
        if prev_snap:
            e["pt"] = round(prev_snap[-1]["value"], 2)
        fm = (e0 or {}).get("fm", 0)
        f = round(flows_on(hist, ds) + fm, 2)
        if f:
            e["f"] = f
        if fm:
            e["fm"] = fm
        if (e0 or {}).get("n"):
            e["n"] = e0["n"]
        x["a"][acc] = e
        got.append(a["name"])
    gaps = sum(1 for y in d["days"] if y["d"] >= "2023-06-16" and "sp" not in y)
    fill_indices(d, rng="max" if gaps > 20 else "10d", today=ds)
    if not x["a"] and "xu" not in x and "sp" not in x:
        d["days"].remove(x)
        print(f"{ds}: piyasa kapalı / veri yok; kayıt açılmadı.")
    for acc in ACCOUNTS:
        recompute(d, acc)
    save(d)
    print(f"{ds}: otomatik kayıt tamam ({len(got)} hesap).")
    return 0


# ------------------------------------------------------------------ IBKR Flex (salt okunur rapor)
FLEX_HOSTS = ("https://ndcdyn.interactivebrokers.com/AccountManagement/FlexWebService",
              "https://gdcdyn.interactivebrokers.com/AccountManagement/FlexWebService")


def _flex_date(s):
    s = (s or "").strip()[:10].replace(";", "")
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s[:8] if fmt == "%Y%m%d" else s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def flex_fetch(token, qid):
    """Flex raporunu XML olarak döndürür. Token/URL asla yazdırılmaz."""
    import time
    import xml.etree.ElementTree as ET
    hdr = {"User-Agent": "portfolio-defter/1.0"}
    ref = url = None
    for host in FLEX_HOSTS:
        try:
            r = requests.get(f"{host}/SendRequest", params={"t": token, "q": qid, "v": 3}, headers=hdr, timeout=30)
            root = ET.fromstring(r.text)
            if root.findtext("Status") == "Success":
                ref, url = root.findtext("ReferenceCode"), root.findtext("Url")
                break
            raise ValueError(f"IBKR Flex isteği reddedildi: {root.findtext('ErrorCode')} {root.findtext('ErrorMessage')}")
        except ValueError:
            raise
        except Exception as ex:
            DIAG.append(f"ibkr SendRequest: {type(ex).__name__}")
    if not ref:
        raise ValueError("IBKR Flex sunucusuna ulaşılamadı")
    for i in range(10):
        time.sleep(5 + 3 * i)
        r = requests.get(url, params={"t": token, "q": ref, "v": 3}, headers=hdr, timeout=60)
        if "<FlexQueryResponse" in r.text[:500]:
            return ET.fromstring(r.text)
        root = ET.fromstring(r.text)
        code = root.findtext("ErrorCode")
        if code not in ("1019", "1018", "1021"):
            raise ValueError(f"IBKR Flex raporu alınamadı: {code} {root.findtext('ErrorMessage')}")
    raise ValueError("IBKR Flex raporu zamanında hazırlanmadı (1019)")


def flex_parse(root):
    nav, flows = {}, {}
    for el in root.iter("EquitySummaryByReportDateInBase"):
        ds = _flex_date(el.get("reportDate"))
        tot = el.get("total")
        if ds and tot not in (None, ""):
            stock = sum(float(el.get(k) or 0) for k in ("stock", "options", "funds", "bonds", "commodities", "crypto"))
            nav[ds] = {"v": round(float(tot), 2), "h": round(stock, 2)}
    for el in root.iter("CashTransaction"):
        if "deposit" not in (el.get("type") or "").lower():
            continue
        ds = _flex_date(el.get("reportDate") or el.get("dateTime") or el.get("settleDate"))
        if not ds:
            continue
        amt = float(el.get("amount") or 0) * float(el.get("fxRateToBase") or 1)
        flows[ds] = round(flows.get(ds, 0) + amt, 2)
    return nav, flows


def flex_trades(root):
    """Flex 'Trades' ve 'Open Positions' bölümleri (sorguda yoksa None döner).
    det: sembol -> {shares, cost, mark, d} (Pozisyonlar sekmesi; cost=costBasisPrice, mark=markPrice)."""
    trades, pos, det = None, None, None
    lods = {(el.get("levelOfDetail") or "EXECUTION").upper() for el in root.iter("Trade")}
    keep = "EXECUTION" if "EXECUTION" in lods else "ORDER"
    for el in root.iter("Trade"):
        if (el.get("levelOfDetail") or "EXECUTION").upper() != keep:
            continue
        ds = _flex_date(el.get("tradeDate") or el.get("dateTime"))
        qty = float(el.get("quantity") or 0)
        if not ds or not qty:
            continue
        trades = trades or []
        trades.append({"id": el.get("tradeID") or el.get("transactionID") or f"{ds}{el.get('symbol')}{qty}",
                       "d": ds, "act": "buy" if qty > 0 else "sell", "sym": el.get("symbol"),
                       "qty": abs(qty), "px": float(el.get("tradePrice") or 0) or None,
                       "amt": abs(float(el.get("tradeMoney") or el.get("proceeds") or 0)) or None,
                       "fee": abs(float(el.get("ibCommission") or 0)) or None})
    for el in root.iter("OpenPosition"):
        if (el.get("levelOfDetail") or "SUMMARY").upper() != "SUMMARY":
            continue
        pos, det = pos or {}, det or {}
        sym = el.get("symbol")
        pos[sym] = float(el.get("position") or 0)
        num = lambda k: (float(el.get(k)) if (el.get(k) or "").strip() else None)
        det[sym] = {"shares": pos[sym], "cost": num("costBasisPrice"), "mark": num("markPrice"),
                    "d": _flex_date(el.get("reportDate"))}
    if pos is None and any(True for _ in root.iter("OpenPositions")):
        pos, det = {}, {}
    return trades, pos, det


def cmd_ibkr(_):
    token, qid = os.environ.get("IBKR_FLEX_TOKEN", "").strip(), os.environ.get("IBKR_FLEX_QUERY_ID", "").strip()
    if not token or not qid:
        print("IBKR Flex secret'ları tanımlı değil; atlandı.")
        return 0
    d = load()
    root = flex_fetch(token, qid)
    nav, flows = flex_parse(root)
    trades, positions, pos_det = flex_trades(root)
    if trades:   # rapor dönemi kısa olabilir: eski işlemler silinmez, yeni gelenler eklenir
        known = {t["id"] for t in d.get("ibkr_trades", [])}
        d.setdefault("ibkr_trades", []).extend(t for t in trades if t["id"] not in known)
    if positions is not None:
        d["ibkr_positions"] = positions
        d["ibkr_pos_detail"] = pos_det
    if not nav:
        raise ValueError("IBKR raporunda 'Net Asset Value (NAV) in Base' bölümü yok; Flex Query'ye ekle.")
    start = ACCOUNTS["ibkr"]["start"]
    last_x = max([ds for ds, e in series(d, "ibkr") if e.get("s") == "x"] or [""])
    # Excel döneminin arasına gün eklenmez (Excel değişimleri bir önceki Excel gününe göredir; araya giren gün çift sayım yapar)
    for x in list(d["days"]):
        e = x.get("a", {}).get("ibkr")
        if e and e.get("s") == "a" and x["d"] < last_x:
            del x["a"]["ibkr"]
            if not x["a"] and not x.get("xu") and not x.get("sp") and not x.get("n"):
                d["days"].remove(x)
    new, cmp_lines, warns = 0, [], []
    # Güvenlik: rapor tüm dönemi kapsasa da yalnızca son 3 takvim günü güncellenir. Daha eski günlere
    # yalnızca defterde hiç IBKR kaydı yoksa (boşluk doldurma) yazılır; var olan eski kayıt asla değişmez.
    win = (datetime.strptime(max(nav), "%Y-%m-%d") - timedelta(days=2)).strftime("%Y-%m-%d")
    for ds in sorted(nav):
        if ds < start:
            continue
        cur = day(d, ds, create=False)
        if ds < win and cur and cur.get("a", {}).get("ibkr") and cur["a"]["ibkr"].get("s") != "x":
            continue
        if cur and (cur.get("a", {}).get("ibkr") or {}).get("s") == "m":
            continue  # elle girilen değer korunur
        if ds < last_x and not day(d, ds, create=False):
            continue
        x = day(d, ds)
        old = x["a"].get("ibkr")
        if not old and ds < last_x:
            continue
        if old and old.get("s") == "x":
            diff = nav[ds]["v"] - old["v"]
            if abs(diff) >= 1:
                cmp_lines.append((ds, old["v"], nav[ds]["v"]))
            continue
        prev = [en["v"] for pd, en in series(d, "ibkr") if pd < ds and en.get("v")]
        if prev and not flows.get(ds) and prev[-1] > 0 and abs(nav[ds]["v"] / prev[-1] - 1) > 0.25:
            warns.append(f"{_tr_date(ds)}: ${prev[-1]:,.0f} → ${nav[ds]['v']:,.0f} (%25+ sıçrama, yazılmadı)")
            if not x["a"] and not x.get("xu") and not x.get("sp") and not x.get("n"):
                d["days"].remove(x)
            continue
        e = {"v": nav[ds]["v"], "h": nav[ds]["h"], "s": "a"}
        if flows.get(ds):
            e["f"] = flows[ds]
        if old and old.get("n"):
            e["n"] = old["n"]
        if old != e:
            x["a"]["ibkr"] = e
            new += 1
    recompute(d, "ibkr")
    save(d)
    msg = [f"🏦 <b>IBKR Flex</b>: {new} gün yazıldı · son rapor günü {_tr_date(max(nav))}"]
    # Rapor tüm dönemi kapsar; mesajda yalnızca son 3 takvim günündeki para hareketleri gösterilir
    cut = (datetime.strptime(max(nav), "%Y-%m-%d") - timedelta(days=2)).strftime("%Y-%m-%d")
    recent = [(k, flows[k]) for k in sorted(flows) if k >= cut and flows[k]]
    if recent:
        msg.append("💵 Para hareketi: " + ", ".join(f"{_tr_date(k)} {v:+,.0f}$" for k, v in recent))
    if warns:
        msg.append("⚠️ Şüpheli değer, elle kontrol et: " + "; ".join(warns))
    if cmp_lines and not d.get("ibkr_cmp_sent"):
        big = sorted(cmp_lines, key=lambda t: -abs(t[2] - t[1]))[:6]
        msg.append("🔎 Excel ↔ IBKR (en büyük farklar): " + "; ".join(
            f"{_tr_date(a)} ${b:,.0f}→${c:,.0f}" for a, b, c in big))
        d["ibkr_cmp_sent"] = True
        save(d)
    common.send_telegram("\n".join(msg))
    return 0


# ------------------------------------------------------------------ Info-TR pozisyonları
def info_prices(tickers):
    import bist_data
    f = bist_data.make_fetch(sorted(tickers))
    return {t: (f(t) or {}).get("c") for t in tickers}


def info_value():
    pf = common.load_json(F_INFO)
    if not pf:
        return None
    pos = pf.get("positions", {})
    px = info_prices(list(pos)) if pos else {}
    missing = [t for t, q in px.items() if not q]
    stocks = sum(p["shares"] * (px.get(t) or p.get("last") or p["cost"]) for t, p in pos.items())
    for t, q in px.items():   # son bilinen fiyatı sakla (fiyat alınamazsa yedek)
        if q:
            pos[t]["last"] = round(q, 4)
    common.save_json(F_INFO, pf)
    return {"total": stocks + pf.get("cash", 0), "stocks": stocks, "missing": missing, "pf": pf}


def cmd_info_trade(_):
    env = os.environ.get
    pf = common.load_json(F_INFO)
    if not pf:
        raise ValueError("Info-TR pozisyon dosyası yok")
    act = (env("IN_ACTION") or "").strip()
    sym = (env("IN_SYMBOL") or "").strip().upper()
    ds = (env("IN_DATE") or "").strip() or datetime.now(timezone(timedelta(hours=3))).date().isoformat()
    datetime.strptime(ds, "%Y-%m-%d")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    pos, txs = pf.setdefault("positions", {}), pf.setdefault("transactions", [])
    if act == "undo":
        if not txs:
            raise ValueError("Geri alınacak işlem yok")
        t = txs.pop()
        pf["cash"], pf["positions"] = t["before"]["cash"], t["before"]["positions"]
        common.save_json(F_INFO, pf)
        common.send_telegram(f"↩️ Info-TR: son işlem geri alındı ({t['action']} {t.get('symbol', '')}).")
        return 0
    before = {"cash": pf.get("cash", 0), "positions": json.loads(json.dumps(pos))}
    tx = {"date": ds, "time": now, "action": act, "before": before}
    if act in ("deposit", "withdraw"):
        amt = _amt(env("IN_QTY"))
        if not amt or amt <= 0:
            raise ValueError("Tutar gir")
        pf["cash"] = round(pf.get("cash", 0) + (amt if act == "deposit" else -amt), 2)
        tx["amount"] = amt
        msg = f"💵 Info-TR: ₺{amt:,.2f} {'yatırıldı' if act == 'deposit' else 'çekildi'}. Nakit ₺{pf['cash']:,.2f}"
    elif act in ("buy", "sell"):
        if not re.fullmatch(r"[A-Z0-9]{2,8}", sym):
            raise ValueError("Geçersiz hisse kodu")
        qty, px, com = _amt(env("IN_QTY")), _amt(env("IN_PRICE")), _amt(env("IN_COMMISSION")) or 0
        if not qty or qty <= 0 or not px or px <= 0:
            raise ValueError("Adet ve fiyat gir")
        ref = info_prices([sym]).get(sym)
        if ref and abs(px / ref - 1) > 0.15 and (env("IN_CONFIRM") or "").lower() != "true":
            raise ValueError(f"{sym} fiyatı ₺{px:,.2f}, son fiyat ₺{ref:,.2f} (%15+ fark). Doğruysa 'confirm' ile tekrar gönder.")
        p = pos.get(sym, {"shares": 0, "cost": 0})
        if act == "buy":
            tot = p["shares"] * p["cost"] + qty * px + com
            p["shares"] = qty + p["shares"]
            p["cost"] = round(tot / p["shares"], 4)
            pos[sym] = p
            pf["cash"] = round(pf.get("cash", 0) - qty * px - com, 2)
        else:
            if qty > p["shares"] + 1e-9:
                raise ValueError(f"{sym}: elde {p['shares']:g} lot var")
            p["shares"] = round(p["shares"] - qty, 6)
            pf["cash"] = round(pf.get("cash", 0) + qty * px - com, 2)
            if p["shares"] <= 1e-9:
                pos.pop(sym, None)
        tx.update(symbol=sym, shares=qty, price=px, commission=com)
        verb = "alındı" if act == "buy" else "satıldı"
        msg = f"🧾 Info-TR: {sym} {qty:g} lot × ₺{px:,.2f} {verb}. Nakit ₺{pf['cash']:,.2f}"
    else:
        raise ValueError("İşlem türü: buy / sell / deposit / withdraw / undo")
    txs.append(tx)
    pf["transactions"] = txs[-500:]
    common.save_json(F_INFO, pf)
    if pf["cash"] < 0:
        msg += "\n⚠️ Nakit eksiye düştü; yatırma girmeyi unutma."
    common.send_telegram(msg)
    return 0


# ------------------------------------------------------------------ elle giriş
def _amt(s):
    s = (s or "").strip().replace(" ", "")
    if not s:
        return None
    if re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+(,\d+)?", s):   # 1.234.567,89
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"[+-]?\d+,\d{1,2}", s):               # 1234,5
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    v = float(s)
    if not (v == v and abs(v) < 1e10):
        raise ValueError("geçersiz tutar")
    return v


def _pairs(s):
    out = {}
    for part in re.split(r"[;\n]", s or ""):
        if not part.strip():
            continue
        k, _, v = part.partition(":")
        acc = ALIASES.get(k.strip().lower())
        if not acc:
            raise ValueError(f"Bilinmeyen hesap: {k.strip()}")
        out[acc] = v.strip()
    return out


def cmd_enter(_):
    env = os.environ.get
    d = load()
    ds = (env("IN_DATE") or "").strip() or datetime.now(timezone(timedelta(hours=3))).date().isoformat()
    datetime.strptime(ds, "%Y-%m-%d")
    x = day(d, ds)
    done = []
    for acc, var in (("info_tr", "IN_INFO_TR"), ("ibkr", "IN_IBKR"), ("midas_tr", "IN_MIDAS_TR"),
                     ("midas_usd", "IN_MIDAS_USD")):
        raw = (env(var) or "").strip()
        if not raw:
            continue
        tot, _, hs = raw.partition("/")
        tot = _amt(tot)
        res = d["accounts"][acc].get("reserve", 0) if ACCOUNTS[acc]["cur"] == "TRY" else 0
        old = x["a"].get(acc, {})
        e = {"v": round(tot - res, 2), "s": "m"}
        if hs.strip():
            e["h"] = _amt(hs)
        if res:
            e["r"] = res
        for kk in ("f", "fm", "n"):
            if kk in old:
                e[kk] = old[kk]
        x["a"][acc] = e
        done.append(ACCOUNTS[acc]["name"])
    for acc, v in _pairs(env("IN_FLOWS")).items():
        e = x["a"].get(acc)
        if not e:
            raise ValueError(f"{ACCOUNTS[acc]['name']} için {ds} değeri yok; akışı değerle birlikte gir.")
        e["f"] = round(e.get("f", 0) + _amt(v), 2)
        e["fm"] = round(e.get("fm", 0) + _amt(v), 2)
        if e.get("s") == "x":
            e["s"] = "m"
        done.append(f"{ACCOUNTS[acc]['name']} akış")
    for acc, v in _pairs(env("IN_RESERVE")).items():
        d["accounts"][acc]["reserve"] = _amt(v)
        e = x["a"].get(acc)
        if e and e.get("s") in ("a", "m"):     # aynı gün değeri yeni yedekle yeniden taban
            tot = e["v"] + e.get("r", 0)
            e["r"] = _amt(v); e["v"] = round(tot - e["r"], 2)
        done.append(f"{ACCOUNTS[acc]['name']} yedek")
    note = (env("IN_NOTE") or "").strip()
    if note:
        k, _, rest = note.partition(":")
        acc = ALIASES.get(k.strip().lower())
        if acc and acc in x["a"]:
            x["a"][acc]["n"] = rest.strip()[:300]
        else:
            x["n"] = note[:300]
        done.append("not")
    if not done:
        raise ValueError("Girilecek bir şey yok.")
    fill_indices(d, today=ds)
    for acc in ACCOUNTS:
        recompute(d, acc)
    save(d)
    lines = []
    for acc in ACCOUNTS:
        e = x["a"].get(acc)
        if e:
            lines.append(_line(d, acc, ds, e))
        else:
            lines.append(f"• {ACCOUNTS[acc]['name']}: ⚠️ girilmedi")
    common.send_telegram(f"📒 <b>Defter {_tr_date(ds)}</b> kaydedildi ({', '.join(done)})\n" + "\n".join(lines))
    return 0


def _prev(d, acc, ds):
    p = None
    for dd, e in series(d, acc):
        if dd >= ds:
            break
        p = e
    return p


def _line(d, acc, ds, e):
    cur = "₺" if ACCOUNTS[acc]["cur"] == "TRY" else "$"
    p = _prev(d, acc, ds)
    pct = (e.get("c", 0) / p["v"] * 100) if p and p.get("v") and "c" in e else None
    arr = "▲" if (e.get("c") or 0) > 0 else "▼" if (e.get("c") or 0) < 0 else "•"
    s = f"• {ACCOUNTS[acc]['name']}: {cur}{e['v']:,.0f}"
    if "c" in e:
        s += f"  {arr} {cur}{e['c']:+,.0f}" + (f" ({pct:+.2f}%)" if pct is not None else "")
    return s


def _tr_date(ds):
    y, m, dd = ds.split("-")
    return f"{dd}.{m}.{y}"


# ------------------------------------------------------------------ dönem hesapları
def daily_rows(d):
    """Her gün, her hesap için yüzde değişim (c / önceki v)."""
    prev = {}
    pidx = {"xu": None, "sp": None}
    out = []
    for x in d["days"]:
        row = {"d": x["d"], "a": {}, "n": x.get("n")}
        for acc, e in x.get("a", {}).items():
            if x["d"] < d["accounts"].get(acc, {}).get("start", ""):
                continue
            p = prev.get(acc)
            pct = e["c"] / p["v"] if p and p.get("v") and "c" in e else None
            row["a"][acc] = dict(e, pct=pct, pv=(p or {}).get("v"))
            prev[acc] = e
        for k in ("xu", "sp"):
            if x.get(k):
                row[k] = x[k]
                row[k + "p"] = (x[k] / pidx[k] - 1) if pidx[k] else None
                pidx[k] = x[k]
        out.append(row)
    return out


def period_key(ds, kind):
    y, m, dd = map(int, ds.split("-"))
    if kind == "W":
        iy, iw, _ = date(y, m, dd).isocalendar()
        return f"{iy}-H{iw:02d}"
    if kind == "M":
        return f"{y}-{m:02d}"
    return str(y)


def periods(d, kind):
    rows = daily_rows(d)
    groups = {}
    order = []
    for r in rows:
        k = period_key(r["d"], kind)
        if k not in groups:
            groups[k] = []
            order.append(k)
        groups[k].append(r)
    out = []
    last_idx = {"xu": None, "sp": None}
    for k in order:
        g = groups[k]
        rec = {"k": k, "from": g[0]["d"], "to": g[-1]["d"], "a": {}}
        for acc in ACCOUNTS:
            es = [r["a"][acc] for r in g if acc in r["a"]]
            if not es:
                continue
            gr, chg, n = 1.0, 0.0, 0
            for e in es:
                if e.get("pct") is not None:
                    gr *= 1 + e["pct"]; chg += e["c"]; n += 1
            rec["a"][acc] = {"v": es[-1]["v"], "c": chg, "pct": (gr - 1) if n else None}
        for grp, accs in (("tl", ("midas_tr", "info_tr")), ("usd", ("midas_usd", "ibkr"))):
            gr, chg, n, val = 1.0, 0.0, 0, 0.0
            for r in g:
                num = sum(r["a"][a]["c"] for a in accs if a in r["a"] and r["a"][a].get("pct") is not None)
                den = sum(r["a"][a]["pv"] for a in accs if a in r["a"] and r["a"][a].get("pct") is not None)
                if den:
                    gr *= 1 + num / den; chg += num; n += 1
            val = sum(rec["a"][a]["v"] for a in accs if a in rec["a"])
            if val:
                rec[grp] = {"v": val, "c": chg, "pct": (gr - 1) if n else None}
        for idx in ("xu", "sp"):
            vals = [r[idx] for r in g if r.get(idx)]
            if vals:
                base = last_idx[idx]
                rec[idx] = vals[-1]
                rec[idx + "p"] = (vals[-1] / base - 1) if base else None
                last_idx[idx] = vals[-1]
        out.append(rec)
    return out


# ------------------------------------------------------------------ Excel
# ------------------------------------------------------------------ işlem kaydı (4 hesap) + kontrol
ACT_TR = {"open": "Açılış", "buy": "Alış", "sell": "Satış", "deposit": "Para yatırma", "withdraw": "Para çekme"}
TX_SRC = {  # hesap -> (işlem dosyası, portföy dosyası, pozisyon anahtarı)
    "midas_usd": ("history.json", "portfolio.json", "portfolio"),
    "midas_tr": ("bist/history.json", "bist/portfolio.json", "portfolio"),
    "info_tr": (F_INFO, F_INFO, "positions"),
}


def _ist(ts):
    try:
        return datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=3)))
    except Exception:
        return None


def trade_rows(d):
    """4 hesabın alım/satım/para hareketleri (geri alınanlar hariç), yeni → eski.
    Dönüş: (satırlar, eldeki pozisyonlar {hesap: {hisse: adet}}, uyarılar)."""
    rows, held, warns = [], {}, []
    for acc, (f_tx, f_pf, key) in TX_SRC.items():
        try:
            txs = (common.load_json(f_tx, {}) or {}).get("transactions", [])
            pf = common.load_json(f_pf, {}) or {}
        except Exception as ex:
            warns.append(f"{ACCOUNTS[acc]['name']}: okunamadı ({type(ex).__name__})")
            continue
        held[acc] = {t: float(p.get("shares") or 0) for t, p in (pf.get(key) or {}).items()}
        for t in txs:
            if t.get("action") not in ACT_TR or t.get("undone"):
                continue
            rows.append({"acc": acc, "d": t.get("date"), "act": t["action"], "sym": t.get("ticker") or t.get("symbol"),
                         "qty": t.get("shares"), "px": t.get("price"), "amt": t.get("amount"),
                         "fee": t.get("commission"), "ts": t.get("time")})
    for acc, m in (d.get("openings") or {}).items():
        for sym, o in m.items():
            rows.append({"acc": acc, "d": o["d"], "act": "open", "sym": sym, "qty": o["qty"]})
    ib = d.get("ibkr_trades") or []
    for t in ib:
        rows.append(dict(t, acc="ibkr"))
    if d.get("ibkr_positions") is not None:
        held["ibkr"] = dict(d["ibkr_positions"])
    elif not ib:
        warns.append("IBKR: Flex sorgusunda 'Trades' ve 'Open Positions' bölümü yok; işlemler listelenmedi")
    for r in rows:
        if r.get("amt") is None and r.get("qty") and r.get("px"):
            r["amt"] = round(r["qty"] * r["px"], 2)
    rows.sort(key=lambda r: (r.get("d") or "", r.get("ts") or ""), reverse=True)
    return rows, held, warns


def recon_rows(rows, held):
    """Eldeki adet ↔ işlem kayıtlarından çıkan net adet. Fark = kayda geçmemiş işlem (veya kayıt öncesi açılış)."""
    net = {}
    for r in rows:
        if r["act"] in ("open", "buy", "sell") and r.get("sym") and r.get("qty"):
            k = (r["acc"], r["sym"])
            net[k] = net.get(k, 0) + (-r["qty"] if r["act"] == "sell" else r["qty"])
    out = []
    for acc in held:
        for sym in sorted(set(held[acc]) | {s for a, s in net if a == acc}):
            h, n = held[acc].get(sym, 0), round(net.get((acc, sym), 0), 6)
            if abs(h - n) < 1e-6:
                continue
            out.append((acc, sym, h, n, round(h - n, 6)))
    return out


def trade_snapshot(d):
    rows, held, warns = trade_rows(d)
    for r in rows:
        r.pop("id", None)
        if r.get("ts"):
            t = _ist(r["ts"])
            r["ts"] = t.strftime("%Y-%m-%d %H:%M") if t else None
    return {"at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "rows": rows,
            "recon": [{"acc": a, "sym": s, "held": h, "net": n, "diff": f} for a, s, h, n, f in recon_rows(rows, held)],
            "miss": missing_days(d), "warns": warns}


def cmd_trades(_):
    """Yalnız İşlemler sekmesini yeniler (pozisyon formlarından çağrılır; hata olsa da formu bozmaz)."""
    try:
        save(load())
        print("islemler yenilendi")
    except Exception as ex:
        print(f"::warning::islemler yenilenemedi ({type(ex).__name__})")
    return 0


def cmd_open(_):
    """Bir kerelik açılış kaydı: kayıt başlamadan önce elde olan adetleri (eldeki − işlemlerden net, yalnız pozitif)
    'Açılış' olarak yazar. Açılışı zaten olan hisseye dokunmaz; eksik satış (negatif fark) işaretli kalır."""
    d = load()
    rows, held, warns = trade_rows(d)
    op = d.setdefault("openings", {})
    first = {}
    for r in rows:
        if r["act"] != "open" and r.get("d"):
            first[r["acc"]] = min(first.get(r["acc"], r["d"]), r["d"])
    today = datetime.now(timezone(timedelta(hours=3))).date().isoformat()
    done, left = [], []
    for acc, sym, h, n, diff in recon_rows(rows, held):
        if diff <= 0 or sym in op.get(acc, {}):
            left.append(f"{ACCOUNTS[acc]['name']} {sym} ({diff:+g})")
            continue
        op.setdefault(acc, {})[sym] = {"qty": diff, "d": first.get(acc, today)}
        done.append(f"{ACCOUNTS[acc]['name']} {sym}")
    save(d)
    msg = ["📒 <b>Açılış kaydı</b>: " + (", ".join(done) if done else "yazılacak fark yok")]
    if left:
        msg.append("⚠️ Hâlâ fark var, elle kontrol et (girilmemiş işlem olabilir): " + "; ".join(left))
    msg += [f"⚠️ {w}" for w in warns]
    common.send_telegram("\n".join(msg))
    return 0


def missing_days(d, back=60):
    """Son `back` gün içinde defterde değeri olmayan iş günleri (resmî tatil olabilir)."""
    have = {x["d"]: set(x.get("a", {})) for x in d.get("days", [])}
    end = date.today() - timedelta(days=1)
    out = {}
    for acc, a in ACCOUNTS.items():
        start = max(end - timedelta(days=back), date.fromisoformat(a.get("start", "2000-01-01")))
        miss, cur = [], start
        while cur <= end:
            if cur.weekday() < 5 and acc not in have.get(cur.isoformat(), ()):
                miss.append(cur.strftime("%d.%m"))
            cur += timedelta(days=1)
        out[acc] = miss
    return out


def build_xlsx(d, out_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    H = Font(bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="0B3D20")
    UP, DN = Font(color="0A7A33"), Font(color="C0392B")
    PCT, NUM, NUM2 = "0.00%", "#,##0", "#,##0.00"
    accs = list(ACCOUNTS)
    names = {a: ACCOUNTS[a]["name"] for a in accs}

    def header(ws, cols, widths=None):
        ws.append(cols)
        for i, _ in enumerate(cols, 1):
            c = ws.cell(1, i)
            c.font, c.fill = H, HF
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 13)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32

    def put(ws, r, c, v, fmt=None, signed=False):
        cell = ws.cell(r, c, v)
        if fmt:
            cell.number_format = fmt
        if signed and isinstance(v, (int, float)) and v:
            cell.font = UP if v > 0 else DN
        return cell

    # Günlük
    ws = wb.active
    ws.title = "Günlük"
    cols = ["Tarih", "Gün"]
    for a in accs:
        cols += [f"{names[a]}\nDeğer", f"{names[a]}\nΔ", f"{names[a]}\n%"]
    cols += ["TL Toplam", "TL Δ", "TL %", "USD Toplam", "USD Δ", "USD %", "XU100", "XU100 %", "S&P 500", "S&P %", "Not"]
    header(ws, cols, {1: 11, 2: 6, len(cols): 40})
    rows = daily_rows(d)
    for i, r in enumerate(rows, 2):
        dt = datetime.strptime(r["d"], "%Y-%m-%d")
        put(ws, i, 1, dt.date(), "dd.mm.yyyy")
        put(ws, i, 2, GUN[dt.weekday()])
        c = 3
        notes = []
        for a in accs:
            e = r["a"].get(a)
            if e:
                put(ws, i, c, e["v"], NUM)
                put(ws, i, c + 1, e.get("c"), NUM, True)
                put(ws, i, c + 2, e.get("pct"), PCT, True)
                if e.get("n"):
                    notes.append(f"{names[a]}: {e['n']}")
                if e.get("f"):
                    notes.append(f"{names[a]} para akışı: {e['f']:+,.0f}")
            c += 3
        for accs_g in (("midas_tr", "info_tr"), ("midas_usd", "ibkr")):
            es = [r["a"][a] for a in accs_g if a in r["a"]]
            ok = [e for e in es if e.get("pct") is not None]
            put(ws, i, c, sum(e["v"] for e in es) if es else None, NUM)
            if ok:
                num, den = sum(e["c"] for e in ok), sum(e["pv"] for e in ok)
                put(ws, i, c + 1, num, NUM, True)
                put(ws, i, c + 2, num / den if den else None, PCT, True)
            c += 3
        put(ws, i, c, r.get("xu"), NUM2); put(ws, i, c + 1, r.get("xup"), PCT, True)
        put(ws, i, c + 2, r.get("sp"), NUM2); put(ws, i, c + 3, r.get("spp"), PCT, True)
        if r.get("n"):
            notes.insert(0, r["n"])
        if notes:
            cell = put(ws, i, c + 4, "📝")
            cell.comment = Comment("\n".join(notes), "Defter")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    # Dönemler
    for title, kind in (("Haftalık", "W"), ("Aylık", "M"), ("Yıllık", "Y")):
        ws = wb.create_sheet(title)
        cols = ["Dönem", "Başlangıç", "Bitiş"]
        for a in accs:
            cols += [f"{names[a]}\nDeğer", f"{names[a]}\nΔ", f"{names[a]}\n%"]
        cols += ["TL Toplam", "TL Δ", "TL %", "USD Toplam", "USD Δ", "USD %", "XU100", "XU100 %", "S&P 500", "S&P %"]
        header(ws, cols)
        for i, p in enumerate(reversed(periods(d, kind)), 2):
            put(ws, i, 1, p["k"])
            put(ws, i, 2, datetime.strptime(p["from"], "%Y-%m-%d").date(), "dd.mm.yyyy")
            put(ws, i, 3, datetime.strptime(p["to"], "%Y-%m-%d").date(), "dd.mm.yyyy")
            c = 4
            for a in accs + ["tl", "usd"]:
                e = p["a"].get(a) if a in ACCOUNTS else p.get(a)
                if e:
                    put(ws, i, c, e["v"], NUM); put(ws, i, c + 1, e["c"], NUM, True)
                    put(ws, i, c + 2, e["pct"], PCT, True)
                c += 3
            put(ws, i, c, p.get("xu"), NUM2); put(ws, i, c + 1, p.get("xup"), PCT, True)
            put(ws, i, c + 2, p.get("sp"), NUM2); put(ws, i, c + 3, p.get("spp"), PCT, True)

    # Gün analizi
    ws = wb.create_sheet("Gün Analizi")
    header(ws, ["Gün", "XU100 ort. %", "XU100 yükseliş oranı", "Gün sayısı", "TL portföy ort. %",
                "TL yükseliş oranı", "S&P ort. %", "S&P yükseliş oranı"], {1: 8})
    stats = weekday_stats(d)
    for i, wd in enumerate(range(5), 2):
        s = stats[wd]
        put(ws, i, 1, GUN[wd])
        put(ws, i, 2, s["xu_avg"], "0.000%", True); put(ws, i, 3, s["xu_up"], "0%"); put(ws, i, 4, s["xu_n"])
        put(ws, i, 5, s["tl_avg"], "0.000%", True); put(ws, i, 6, s["tl_up"], "0%")
        put(ws, i, 7, s["sp_avg"], "0.000%", True); put(ws, i, 8, s["sp_up"], "0%")

    # İşlemler: 4 hesabın alım/satım/para hareketi kaydı
    rows, held, warns = trade_rows(d)
    ws = wb.create_sheet("İşlemler", 1)
    cols = ["İşlem tarihi", "Gün", "Hesap", "İşlem", "Hisse", "Adet", "Fiyat", "Tutar", "Komisyon",
            "Para", "Kayda girildi (İst.)", "Gecikme (gün)"]
    header(ws, cols, {3: 11, 4: 13, 11: 17})
    for i, r in enumerate(rows, 2):
        dt = datetime.strptime(r["d"], "%Y-%m-%d").date() if r.get("d") else None
        cur = ACCOUNTS[r["acc"]]["cur"]
        put(ws, i, 1, dt, "dd.mm.yyyy"); put(ws, i, 2, GUN[dt.weekday()] if dt else None)
        put(ws, i, 3, ACCOUNTS[r["acc"]]["name"]); put(ws, i, 4, ACT_TR.get(r["act"], r["act"]))
        put(ws, i, 5, r.get("sym")); put(ws, i, 6, r.get("qty"), "#,##0.####")
        put(ws, i, 7, r.get("px"), "#,##0.00##"); put(ws, i, 8, r.get("amt"), NUM2)
        put(ws, i, 9, r.get("fee"), NUM2); put(ws, i, 10, "₺" if cur == "TRY" else "$")
        ts = _ist(r["ts"]) if r.get("ts") else None
        if ts:
            put(ws, i, 11, ts.replace(tzinfo=None), "dd.mm.yyyy hh:mm")
            if dt:
                lag = (ts.date() - dt).days
                put(ws, i, 12, lag)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    # Kontrol: eksik kayıt ipuçları
    ws = wb.create_sheet("Kontrol", 2)
    header(ws, ["Hesap", "Hisse", "Eldeki adet", "İşlemlerden net adet", "Fark", "Olası neden"], {6: 48})
    i = 2
    for acc, sym, h, n, diff in recon_rows(rows, held):
        why = ("Kayıt başlamadan önce açılmış pozisyon ya da girilmemiş alış" if diff > 0
               else "Girilmemiş satış ya da kayıt öncesi alış")
        for c, v in enumerate((ACCOUNTS[acc]["name"], sym, h, n, diff, why), 1):
            put(ws, i, c, v, "#,##0.####" if c in (3, 4, 5) else None)
        ws.cell(i, 5).font = DN
        i += 1
    if i == 2:
        put(ws, i, 1, "✅ Eldeki adetler işlem kayıtlarıyla tutarlı"); i += 1
    i += 1
    put(ws, i, 1, "Son 60 günde defter kaydı olmayan iş günleri (tatil olabilir)").font = Font(bold=True)
    for acc, miss in missing_days(d).items():
        i += 1
        put(ws, i, 1, ACCOUNTS[acc]["name"]); put(ws, i, 2, ", ".join(miss) if miss else "✅ eksik yok")
    for w in warns:
        i += 2
        put(ws, i, 1, f"⚠️ {w}")
    wb.save(out_path)


def weekday_stats(d, since=None):
    rows = daily_rows(d)
    st = {i: {} for i in range(7)}
    acc = {i: {"xu": [], "sp": [], "tl": []} for i in range(7)}
    for r in rows:
        if since and r["d"] < since:
            continue
        wd = datetime.strptime(r["d"], "%Y-%m-%d").weekday()
        if r.get("xup") is not None:
            acc[wd]["xu"].append(r["xup"])
        if r.get("spp") is not None:
            acc[wd]["sp"].append(r["spp"])
        ok = [r["a"][a] for a in ("midas_tr", "info_tr") if a in r["a"] and r["a"][a].get("pct") is not None]
        den = sum(e["pv"] for e in ok)
        if den:
            acc[wd]["tl"].append(sum(e["c"] for e in ok) / den)
    for wd in range(7):
        for k in ("xu", "sp", "tl"):
            v = acc[wd][k]
            st[wd][k + "_avg"] = sum(v) / len(v) if v else None
            st[wd][k + "_up"] = sum(1 for x in v if x > 0) / len(v) if v else None
            st[wd][k + "_n"] = len(v)
    return st


def send_document(path, caption):
    token, chat = os.environ.get("TELEGRAM_BOT_TOKEN"), os.environ.get("TELEGRAM_CHAT_ID")
    if not token or not chat:
        print("[WARN] Telegram bilgileri yok; dosya gönderilmedi.")
        return False
    with open(path, "rb") as f:
        r = requests.post(f"https://api.telegram.org/bot{token}/sendDocument",
                          data={"chat_id": chat, "caption": caption[:1000], "parse_mode": "HTML"},
                          files={"document": (Path(path).name, f)}, timeout=60)
    if r.status_code != 200:
        print(f"[WARN] Telegram sendDocument: HTTP {r.status_code}")
    return r.status_code == 200


def week_summary(d):
    ps = periods(d, "W")
    if not ps:
        return ""
    p = ps[-1]
    L = [f"📒 <b>Haftalık özet</b> {_tr_date(p['from'])} – {_tr_date(p['to'])}"]

    def fmt(e, cur):
        if not e or e.get("pct") is None:
            return "—"
        arr = "▲" if e["c"] > 0 else "▼" if e["c"] < 0 else "•"
        return f"{arr} {e['pct'] * 100:+.2f}% ({cur}{e['c']:+,.0f})"
    for a in ACCOUNTS:
        cur = "₺" if ACCOUNTS[a]["cur"] == "TRY" else "$"
        L.append(f"• {ACCOUNTS[a]['name']}: {fmt(p['a'].get(a), cur)}")
    L.append(f"• <b>TL toplam</b>: {fmt(p.get('tl'), '₺')} · XU100 "
             + (f"{p['xup'] * 100:+.2f}%" if p.get("xup") is not None else "—"))
    L.append(f"• <b>USD toplam</b>: {fmt(p.get('usd'), '$')} · S&P 500 "
             + (f"{p['spp'] * 100:+.2f}%" if p.get("spp") is not None else "—"))
    miss = [f"{_tr_date(r['d'])} {ACCOUNTS[a]['name']}" for r in d["days"] if p["from"] <= r["d"] <= p["to"]
            for a in ("info_tr", "ibkr") if a not in r.get("a", {}) and (r.get("xu") or r.get("sp"))
            and r["d"] >= ACCOUNTS[a].get("start", "")]
    if miss:
        L.append("⚠️ Girilmemiş: " + ", ".join(miss[:12]))
    return "\n".join(L)


def cmd_excel(args):
    d = load()
    out = Path(os.environ.get("RUNNER_TEMP", "/tmp")) / f"portfoy_defteri_{date.today():%Y%m%d}.xlsx"
    build_xlsx(d, out)
    cap = week_summary(d) if args.summary else "📒 Portföy Defteri (Excel)"
    send_document(out, cap)
    print("Excel üretildi.")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser()
    sp = ap.add_subparsers(dest="cmd", required=True)
    sp.add_parser("kutu-anahtar").set_defaults(fn=cmd_box_keygen)
    sp.add_parser("ice-aktar").set_defaults(fn=cmd_import)
    a = sp.add_parser("oto"); a.add_argument("--date"); a.set_defaults(fn=cmd_auto)
    sp.add_parser("gir").set_defaults(fn=cmd_enter)
    sp.add_parser("ibkr").set_defaults(fn=cmd_ibkr)
    sp.add_parser("acilis").set_defaults(fn=cmd_open)
    sp.add_parser("islemler").set_defaults(fn=cmd_trades)
    sp.add_parser("info-islem").set_defaults(fn=cmd_info_trade)
    e = sp.add_parser("excel"); e.add_argument("--summary", action="store_true"); e.set_defaults(fn=cmd_excel)
    args = ap.parse_args(argv)
    try:
        return args.fn(args)
    except (ValueError, common.VaultError) as ex:
        common.send_telegram(f"❌ Defter: {ex}")
        print(f"::error::{type(ex).__name__}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
